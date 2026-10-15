from __future__ import annotations

import hashlib
import io
import re
import unicodedata
from decimal import Decimal, InvalidOperation
from pathlib import Path

import pandas as pd
import streamlit as st
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

APP_VERSION = "0.1.3"
MIRO_SCORE_COLUMN = 8  # H: Moodle grade
FIRST_WEEK_COLUMN = 10  # J: keep I (last downloaded) intact
SETTINGS_SHEET = "Miro_Settings"
DEFAULT_WEEKS = 14
DEFAULT_WEEK_MAX = 2.0
DEFAULT_FINAL_MAX = 20.0
DEFAULT_DECIMALS = 1


def normalize_number(value):
    if value is None or value == "":
        return None
    try:
        n = float(value)
        if n.is_integer():
            return int(n)
        return n
    except (TypeError, ValueError):
        return value


def normalized_id_text(value) -> str:
    """Render numbers and text IDs consistently, including Excel-style '1234.0'."""
    if value is None or value == "":
        return ""
    if isinstance(value, bool):
        return ""
    if isinstance(value, (int, float)):
        if isinstance(value, float) and value.is_integer():
            return str(int(value))
        return str(value)

    raw = unicodedata.normalize("NFKC", str(value)).strip()
    if not raw or raw.startswith("="):
        # Do not accidentally identify the cell reference in an Excel formula.
        return ""
    # Sometimes a CSV import or Excel reformatting produces an integer ID with
    # a trailing .0, or the string "1.2345678E+07".
    try:
        if re.fullmatch(r"[+\-]?[0-9]+(?:\.0+)?", raw) or re.fullmatch(
            r"[+\-]?[0-9]+(?:\.[0-9]+)?[Ee][+\-]?[0-9]+", raw
        ):
            number = Decimal(raw)
            if number == number.to_integral_value():
                return format(number, "f").split(".", 1)[0]
    except InvalidOperation:
        pass
    return raw


def id_last4(value) -> str:
    normalized = normalized_id_text(value)
    digits = "".join(re.findall(r"[0-9]", normalized))
    if digits:
        return digits[-4:].zfill(4)
    return ""


def find_header_row(ws) -> int | None:
    for r in range(1, min(ws.max_row, 12) + 1):
        b = str(ws.cell(r, 2).value or "").strip()
        c = str(ws.cell(r, 3).value or "").strip()
        h = str(ws.cell(r, MIRO_SCORE_COLUMN).value or "").strip()
        if b == "氏名" and "ID" in c.upper() and "Miro" in h:
            return r
    return None


def candidate_sheets(wb):
    result = []
    for ws in wb.worksheets:
        if ws.title == SETTINGS_SHEET:
            continue
        row = find_header_row(ws)
        if row:
            result.append((ws.title, row))
    return result


def ensure_settings_sheet(wb):
    if SETTINGS_SHEET in wb.sheetnames:
        ws = wb[SETTINGS_SHEET]
    else:
        ws = wb.create_sheet(SETTINGS_SHEET)
        ws["A1"] = "Miro Grader settings"
        ws["A2"] = "app_version"
        ws["B2"] = APP_VERSION
        ws["A3"] = "num_weeks"
        ws["B3"] = DEFAULT_WEEKS
        ws["A4"] = "week_max"
        ws["B4"] = DEFAULT_WEEK_MAX
        ws["A5"] = "final_max"
        ws["B5"] = DEFAULT_FINAL_MAX
        ws["A6"] = "decimals"
        ws["B6"] = DEFAULT_DECIMALS
        ws["A7"] = "current_week"
        ws["B7"] = 1
        ws["A8"] = "active_weeks"
        ws["B8"] = ""
    ws.sheet_state = "hidden"
    return ws


def read_settings(wb):
    ws = ensure_settings_sheet(wb)
    settings = {
        "num_weeks": int(ws["B3"].value or DEFAULT_WEEKS),
        "week_max": float(ws["B4"].value or DEFAULT_WEEK_MAX),
        "final_max": float(ws["B5"].value or DEFAULT_FINAL_MAX),
        "decimals": int(ws["B6"].value if ws["B6"].value is not None else DEFAULT_DECIMALS),
        "current_week": int(ws["B7"].value or 1),
        "active_weeks": [],
    }
    active_raw = str(ws["B8"].value or "").strip()
    if active_raw:
        settings["active_weeks"] = sorted({int(x) for x in active_raw.split(",") if x.strip().isdigit()})
    settings["num_weeks"] = max(1, min(40, settings["num_weeks"]))
    settings["current_week"] = max(1, min(settings["num_weeks"], settings["current_week"]))
    settings["active_weeks"] = [w for w in settings["active_weeks"] if 1 <= w <= settings["num_weeks"]]
    return settings


def write_settings(wb, settings):
    ws = ensure_settings_sheet(wb)
    ws["B2"] = APP_VERSION
    ws["B3"] = int(settings["num_weeks"])
    ws["B4"] = float(settings["week_max"])
    ws["B5"] = float(settings["final_max"])
    ws["B6"] = int(settings["decimals"])
    ws["B7"] = int(settings["current_week"])
    ws["B8"] = ",".join(str(x) for x in sorted(settings["active_weeks"]))
    ws.sheet_state = "hidden"


def week_header(week_no: int) -> str:
    return f"Week {week_no}"


def week_column_map(ws, header_row: int):
    mapping = {}
    for col in range(FIRST_WEEK_COLUMN, ws.max_column + 1):
        value = str(ws.cell(header_row, col).value or "").strip()
        m = re.fullmatch(r"Week\s*(\d+)", value, flags=re.I)
        if m:
            mapping[int(m.group(1))] = col
    return mapping


def ensure_week_columns(ws, header_row: int, num_weeks: int):
    mapping = week_column_map(ws, header_row)
    next_col = max(FIRST_WEEK_COLUMN, ws.max_column + 1)
    for week in range(1, num_weeks + 1):
        if week not in mapping:
            col = next_col
            ws.cell(header_row, col).value = week_header(week)
            ws.cell(header_row, col).font = Font(bold=True)
            ws.cell(header_row, col).alignment = Alignment(horizontal="center")
            ws.cell(header_row, col).fill = PatternFill(fill_type="solid", fgColor="E8F0FE")
            mapping[week] = col
            next_col += 1
    return mapping


def student_rows(ws, header_row: int):
    """Read English names (A), Japanese names (B), and student IDs (C)."""
    rows = []
    for r in range(header_row + 1, ws.max_row + 1):
        name_en = ws.cell(r, 1).value
        name = ws.cell(r, 2).value
        student_id = ws.cell(r, 3).value
        if all(value is None or not str(value).strip() for value in (name_en, name, student_id)):
            continue
        rows.append(
            {
                "row": r,
                "id4": id_last4(student_id),
                "name": str(name or "").strip(),
                "name_en": str(name_en or "").strip(),
                "id_full": student_id,
            }
        )
    return rows


def normalize_search(value) -> str:
    """Normalize case, full-width characters and whitespace for matching."""
    raw = unicodedata.normalize("NFKC", str(value or "")).casefold()
    return re.sub(r"[\s\u200b\u3000・･.,_\-]+", "", raw)


def student_display(student) -> str:
    japanese = student["name"]
    english = student.get("name_en", "")
    names = f"{japanese} / {english}" if japanese and english else japanese or english or "（氏名なし）"
    return f"{student['id4'] or '----'} | {names}"


def student_match_score(student, query: str) -> int:
    """Find candidates even from a longer copied Miro sticky-note body."""
    query = str(query or "").strip()
    q = normalize_search(query)
    if not q:
        return 0

    digits_in_query = "".join(re.findall(r"[0-9]", unicodedata.normalize("NFKC", query)))
    last4 = student["id4"]
    full_id = normalize_search(normalized_id_text(student.get("id_full")))
    if last4 and len(last4) == 4 and last4 in digits_in_query:
        return 120
    if full_id and len(full_id) >= 4 and q in full_id:
        return 110

    for name in (student["name"], student.get("name_en", "")):
        normalized = normalize_search(name)
        if not normalized:
            continue
        if q == normalized:
            return 100
        if q in normalized:
            return 90
        if len(normalized) >= 3 and normalized in q:
            return 80
        # English given-name/family-name order may be reversed in Miro.
        parts = [normalize_search(x) for x in re.split(r"[\s,]+", name.strip())]
        parts = [x for x in parts if len(x) >= 2]
        if len(parts) >= 2 and all(x in q for x in parts):
            return 70
    return 0


def compute_raw_total(ws, row: int, week_cols: dict[int, int], active_weeks: list[int]) -> float:
    total = 0.0
    for week in active_weeks:
        col = week_cols.get(week)
        if not col:
            continue
        value = ws.cell(row, col).value
        try:
            if value is not None and value != "":
                total += float(value)
        except (TypeError, ValueError):
            pass
    return total


def recalc_final_scores(ws, header_row: int, settings, week_cols):
    active_weeks = settings["active_weeks"]
    week_max = float(settings["week_max"])
    final_max = float(settings["final_max"])
    decimals = int(settings["decimals"])
    possible = len(active_weeks) * week_max

    for stu in student_rows(ws, header_row):
        row = stu["row"]
        if possible <= 0:
            ws.cell(row, MIRO_SCORE_COLUMN).value = None
            continue
        raw = compute_raw_total(ws, row, week_cols, active_weeks)
        final = raw / possible * final_max
        final = max(0.0, min(final_max, final))
        rounded = round(final, decimals)
        ws.cell(row, MIRO_SCORE_COLUMN).value = int(rounded) if decimals == 0 else rounded


def workbook_bytes(wb) -> bytes:
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def build_table(ws, header_row: int, settings, week_cols, current_week: int):
    data = []
    active = settings["active_weeks"]
    possible = len(active) * settings["week_max"]
    current_col = week_cols[current_week]
    for stu in student_rows(ws, header_row):
        row = stu["row"]
        raw = compute_raw_total(ws, row, week_cols, active)
        final = None
        if possible > 0:
            final = round(raw / possible * settings["final_max"], settings["decimals"])
        data.append(
            {
                "下4桁": stu["id4"],
                "氏名": stu["name"],
                week_header(current_week): normalize_number(ws.cell(row, current_col).value),
                "素点合計": round(raw, 2),
                f"最終 / {settings['final_max']:g}": final,
            }
        )
    return pd.DataFrame(data)


st.set_page_config(page_title="Miro Grader", page_icon="🧩", layout="wide")
st.title("🧩 Miro Group Work Grader")
st.caption(f"アプリバージョン: v{APP_VERSION}")
st.caption("ExcelのJ列以降にWeek別点数を保存し、H列へMiroの最終評定を自動反映します。I列は保持します。")

uploaded = st.file_uploader("評定Excel（.xlsx）を選択", type=["xlsx"])
if uploaded is None:
    st.info("まず評定Excelを選択してください。元ファイルは上書きせず、更新版をダウンロードできます。")
    st.stop()

file_data = uploaded.getvalue()
file_hash = hashlib.sha256(file_data).hexdigest()
if st.session_state.get("loaded_hash") != file_hash:
    try:
        wb = load_workbook(io.BytesIO(file_data))
    except Exception as e:
        st.error(f"Excelを開けませんでした: {e}")
        st.stop()
    candidates = candidate_sheets(wb)
    if not candidates:
        st.error("B列「氏名」、C列「IDナンバ」、H列「グループワーク (Miro)」の見出しを持つシートが見つかりませんでした。")
        st.stop()
    st.session_state.pop("selected_student_rows", None)
    st.session_state.loaded_hash = file_hash
    st.session_state.wb = wb
    st.session_state.target_sheet = candidates[0][0]
    st.session_state.header_row = candidates[0][1]
    st.session_state.original_name = uploaded.name
    st.session_state.flash = "Excelを読み込みました。"

wb = st.session_state.wb
available_sheets = [ws.title for ws in wb.worksheets if ws.title != SETTINGS_SHEET]
selected_sheet = st.sidebar.selectbox(
    "評定シート",
    available_sheets,
    index=available_sheets.index(st.session_state.target_sheet) if st.session_state.target_sheet in available_sheets else 0,
)
if selected_sheet != st.session_state.target_sheet:
    ws_test = wb[selected_sheet]
    detected = find_header_row(ws_test)
    if not detected:
        st.sidebar.error("このシートでは見出し行を自動認識できません。")
        st.stop()
    st.session_state.target_sheet = selected_sheet
    st.session_state.header_row = detected
    st.session_state.pop("selected_student_rows", None)

ws = wb[st.session_state.target_sheet]
header_row = st.session_state.header_row
settings = read_settings(wb)

st.sidebar.subheader("設定")
with st.sidebar.form("settings_form"):
    num_weeks = st.number_input("Week数", min_value=1, max_value=40, value=int(settings["num_weeks"]), step=1)
    week_max = st.number_input("1回あたりの満点", min_value=0.1, max_value=100.0, value=float(settings["week_max"]), step=0.5)
    final_max = st.number_input("Miro最終配点（/100）", min_value=0.1, max_value=100.0, value=float(settings["final_max"]), step=1.0)
    decimals = st.number_input("最終点の小数桁", min_value=0, max_value=3, value=int(settings["decimals"]), step=1)
    all_weeks_for_form = list(range(1, int(num_weeks) + 1))
    active_default = [w for w in settings["active_weeks"] if w in all_weeks_for_form]
    active_weeks = st.multiselect(
        "採点対象Week（飛ばした週は外す）",
        all_weeks_for_form,
        default=active_default,
        format_func=lambda x: f"Week {x}",
    )
    save_settings = st.form_submit_button("設定を保存", use_container_width=True)

if save_settings:
    settings["num_weeks"] = int(num_weeks)
    settings["week_max"] = float(week_max)
    settings["final_max"] = float(final_max)
    settings["decimals"] = int(decimals)
    settings["active_weeks"] = sorted(active_weeks)
    settings["current_week"] = min(settings["current_week"], settings["num_weeks"])
    week_cols = ensure_week_columns(ws, header_row, settings["num_weeks"])
    recalc_final_scores(ws, header_row, settings, week_cols)
    write_settings(wb, settings)
    st.session_state.flash = "設定を保存し、H列を再計算しました。"
    st.rerun()

# Ensure columns exist based on current settings.
week_cols = ensure_week_columns(ws, header_row, settings["num_weeks"])
roster = student_rows(ws, header_row)

# Show what was actually read from the workbook. This makes it possible to
# distinguish a search problem from a missing/misformatted ID column.
with st.expander("🔎 名簿の読み取り確認（下4桁で候補が出ない場合）", expanded=False):
    parsed = sum(bool(s["id4"]) for s in roster)
    st.write(
        f"対象シート: **{ws.title}** ／見出し行: **{header_row}行目** ／ "
        f"学生: **{len(roster)}人** ／C列から下4桁を抽出できた学生: **{parsed}人**"
    )
    if parsed == 0:
        st.error("C列から学籍番号の数字を読み取れていません。見出しと実際のID列を確認してください。")
    elif parsed < len(roster):
        st.warning("一部の学生でC列のIDが空欄、数式、または数字なしになっています。")
    if roster:
        st.dataframe(
            pd.DataFrame(
                {
                    "Excel行": s["row"],
                    "A列（英語名）": s["name_en"],
                    "B列（氏名）": s["name"],
                    "C列（Excel上の値）": str(s["id_full"] or ""),
                    "検索に使う下4桁": s["id4"] or "（抽出できず）",
                }
                for s in roster
            ),
            hide_index=True,
            use_container_width=True,
            height=245,
        )
        st.caption("ここに表示された下4桁が検索可能な番号です。表示が違う場合はExcelのC列を確認してください。")

if st.session_state.get("flash"):
    st.success(st.session_state.pop("flash"))

c1, c2, c3, c4 = st.columns(4)
c1.metric("学生数", len(roster))
c2.metric("1回満点", f"{settings['week_max']:g}")
c3.metric("採点対象Week", len(settings["active_weeks"]))
c4.metric("最終配点", f"{settings['final_max']:g} / 100")

week_options = list(range(1, settings["num_weeks"] + 1))
default_index = max(0, min(len(week_options) - 1, settings["current_week"] - 1))
current_week = st.selectbox("今回のWeek", week_options, index=default_index, format_func=lambda x: f"Week {x}")
settings["current_week"] = current_week
write_settings(wb, settings)

current_col = week_cols[current_week]
entered_count = sum(1 for stu in roster if ws.cell(stu["row"], current_col).value not in (None, ""))
st.progress(entered_count / len(roster) if roster else 0, text=f"Week {current_week}: {entered_count} / {len(roster)} 人入力済み")

st.subheader(f"Week {current_week} の採点")
st.caption("A列の英語名・B列の氏名・C列の学籍番号（下4桁）から候補を検索します。")
search = st.text_area(
    "学生の候補を探す（Miroの付箋本文も貼り付け可能）",
    placeholder="例: 1234 / 山田 / Taro Yamada / Miroの投稿内容",
    height=80,
)

# Keep the group selection independent of the search query. Filtering a
# multiselect's options on every keystroke would discard earlier selections.
roster_by_row = {s["row"]: s for s in roster}
if "selected_student_rows" not in st.session_state:
    st.session_state["selected_student_rows"] = []


def add_candidate(row_number: int):
    selected = list(st.session_state.get("selected_student_rows", []))
    if row_number not in selected:
        selected.append(row_number)
        st.session_state["selected_student_rows"] = selected


if search.strip():
    candidates = sorted(
        ((student_match_score(s, search), s) for s in roster),
        key=lambda item: (-item[0], item[1]["row"]),
    )
    candidates = [s for score, s in candidates if score > 0]
    st.caption(f"該当候補: {len(candidates)} 人（候補をクリックするとグループに追加）")
    for i in range(0, min(len(candidates), 12), 2):
        cols = st.columns(2)
        for col, stu in zip(cols, candidates[i : i + 2]):
            with col:
                st.button(
                    f"＋ {student_display(stu)}",
                    key=f"candidate_add_{stu['row']}",
                    on_click=add_candidate,
                    args=(stu["row"],),
                    use_container_width=True,
                )
    if not candidates:
        st.warning("該当者がいません。上の『名簿の読み取り確認』を開き、ExcelのC列と抽出された下4桁を確認してください。下の選択欄からも探せます。")

selected_rows = st.multiselect(
    "このグループの学生を選択（直接検索も可能）",
    options=list(roster_by_row.keys()),
    format_func=lambda row_number: student_display(roster_by_row[row_number]),
    key="selected_student_rows",
    placeholder="氏名・英語名・下4桁を入力して検索・追加",
)
selected_students = [roster_by_row[r] for r in selected_rows]

score_col, action_col = st.columns([1, 2])
with score_col:
    score = st.number_input(
        "点数",
        min_value=0.0,
        max_value=float(settings["week_max"]),
        value=float(settings["week_max"]),
        step=0.5 if settings["week_max"] % 1 else 1.0,
    )
with action_col:
    st.write(" ")
    st.write(" ")
    apply_score = st.button(
        f"選択した {len(selected_students)} 人に {score:g} 点を入力",
        type="primary",
        disabled=not selected_students,
        use_container_width=True,
    )

# Quick score buttons for small integer scales.
if float(settings["week_max"]).is_integer() and int(settings["week_max"]) <= 5:
    quick_values = list(range(0, int(settings["week_max"]) + 1))
    quick_cols = st.columns(len(quick_values))
    quick_clicked = None
    for idx, value in enumerate(quick_values):
        if quick_cols[idx].button(f"{value}点", disabled=not selected_students, use_container_width=True, key=f"q_{value}"):
            quick_clicked = value
    if quick_clicked is not None:
        score = float(quick_clicked)
        apply_score = True

if apply_score:
    for stu in selected_students:
        ws.cell(stu["row"], current_col).value = normalize_number(score)
    if current_week not in settings["active_weeks"]:
        settings["active_weeks"].append(current_week)
        settings["active_weeks"] = sorted(set(settings["active_weeks"]))
    settings["current_week"] = current_week
    recalc_final_scores(ws, header_row, settings, week_cols)
    write_settings(wb, settings)
    st.session_state.flash = f"Week {current_week}: {len(selected_students)}人に {score:g} 点を入力しました。"
    st.rerun()

clear_score = st.button("選択した学生のこのWeekの点数を空欄に戻す", disabled=not selected_students)
if clear_score:
    for stu in selected_students:
        ws.cell(stu["row"], current_col).value = None
    recalc_final_scores(ws, header_row, settings, week_cols)
    write_settings(wb, settings)
    st.session_state.flash = f"Week {current_week}: {len(selected_students)}人の点数を空欄に戻しました。"
    st.rerun()

st.divider()
st.subheader("確認")
df = build_table(ws, header_row, settings, week_cols, current_week)
st.dataframe(df, use_container_width=True, hide_index=True, height=min(700, 80 + 35 * max(1, len(df))))

if settings["active_weeks"]:
    st.caption(
        "H列は、採点対象Weekの素点合計 ÷ 採点対象Weekの満点合計 × Miro最終配点、で計算します。"
        " 採点対象Week内の空欄は、その学生について0点相当になります。"
    )
else:
    st.warning("採点対象Weekがまだありません。最初の採点を入力すると、そのWeekが自動で採点対象になります。")

# Recalculate once more before export.
recalc_final_scores(ws, header_row, settings, week_cols)
write_settings(wb, settings)

original = Path(st.session_state.original_name)
out_name = f"{original.stem}_Miro更新{original.suffix}"
export_data = workbook_bytes(wb)
st.download_button(
    "更新済みExcelをダウンロード",
    data=export_data,
    file_name=out_name,
    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    type="primary",
    use_container_width=True,
)
st.caption("元ファイルは上書きしません。次回は、ここからダウンロードした更新版Excelを読み込んで続けてください。")
